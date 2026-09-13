import logging
from pathlib import Path

logger = logging.getLogger(__name__)

def extract_text(filepath: Path, filetype: str) -> str:
    """
    Extracts text from a given file based on its file type.
    
    Args:
        filepath (Path): The path to the file.
        filetype (str): The extension or mime-type representing the file type.
        
    Returns:
        str: The extracted text, or empty string on failure.
    """
    ext = filetype.lower()
    if not ext.startswith('.'):
        ext = '.' + ext.split('/')[-1] if '/' in ext else '.' + ext

    text = ""
    try:
        if ext in ['.pdf']:
            try:
                import pdfplumber
                with pdfplumber.open(filepath) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
            except ImportError:
                logger.warning("pdfplumber not installed. Cannot extract PDF text.")
                
        elif ext in ['.docx']:
            try:
                import docx
                doc = docx.Document(filepath)
                text = "\n".join([para.text for para in doc.paragraphs])
            except ImportError:
                logger.warning("python-docx not installed. Cannot extract DOCX text.")
                
        elif ext in ['.txt', '.csv']:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
                
        elif ext in ['.png', '.jpg', '.jpeg', '.tiff', '.bmp']:
            try:
                from PIL import Image
                import pytesseract
                try:
                    from app.config import TESSERACT_CMD
                    if TESSERACT_CMD:
                        pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
                except ImportError:
                    pass
                text = pytesseract.image_to_string(Image.open(filepath))
            except ImportError:
                logger.warning("Pillow or pytesseract not installed. Cannot extract image text.")
            except pytesseract.TesseractNotFoundError:
                logger.error("SYSTEM ERROR: Tesseract is not installed on this machine. OCR will fail.")
            except Exception as e:
                logger.warning(f"OCR failed for {filepath}: {e}")
                
        elif ext in ['.xlsx']:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(filepath, data_only=True)
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " ".join([str(cell) for cell in row if cell is not None])
                        if row_text.strip():
                            text += row_text + "\n"
            except ImportError:
                logger.warning("openpyxl not installed. Cannot extract XLSX text.")
            except Exception as e:
                logger.warning(f"Failed to extract text from XLSX {filepath}: {e}")
                
        elif ext in ['.xls']:
            try:
                import xlrd
                wb = xlrd.open_workbook(filepath)
                for sheet in wb.sheets():
                    for rx in range(sheet.nrows):
                        row_text = " ".join([str(cell) for cell in sheet.row_values(rx) if cell])
                        if row_text.strip():
                            text += row_text + "\n"
            except ImportError:
                logger.warning("xlrd not installed. Cannot extract XLS text.")
            except Exception as e:
                logger.warning(f"Failed to extract text from XLS {filepath}: {e}")
                
        elif ext in ['.eml']:
            try:
                import email
                from email import policy
                with open(filepath, 'rb') as f:
                    msg = email.message_from_binary_file(f, policy=policy.default)
                
                text += f"Subject: {msg.get('subject', '')}\n"
                text += f"From: {msg.get('from', '')}\n"
                text += f"To: {msg.get('to', '')}\n\n"
                
                body = msg.get_body(preferencelist=('plain', 'html'))
                if body:
                    text += body.get_content()
            except Exception as e:
                logger.warning(f"Failed to extract text from EML {filepath}: {e}")
                
        elif ext in ['.msg']:
            try:
                import extract_msg
                msg = extract_msg.Message(filepath)
                text += f"Subject: {msg.subject}\n"
                text += f"From: {msg.sender}\n"
                text += f"To: {msg.to}\n\n"
                text += msg.body if msg.body else ""
                msg.close()
            except ImportError:
                logger.warning("extract-msg not installed. Cannot extract MSG text.")
            except Exception as e:
                logger.warning(f"Failed to extract text from MSG {filepath}: {e}")
                
                
        else:
            logger.info(f"No extraction logic for filetype {ext}")
            
    except Exception as e:
        logger.error(f"Error extracting text from {filepath}: {e}")
        
    return text.strip()
