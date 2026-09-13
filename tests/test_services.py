import pytest
from app.services.tagger import generate_tags
from app.services.extractor import extract_text

def test_generate_tags():
    text = "The quick brown fox jumps over the lazy dog. The fox is very quick."
    tags = generate_tags(text)
    
    assert isinstance(tags, list)
    assert len(tags) <= 10  # Configured MAX_AUTO_TAGS
    
    # Should extract meaningful phrases, often includes 'quick' or 'fox'
    assert any("fox" in tag.lower() for tag in tags) or any("quick" in tag.lower() for tag in tags)

def test_extract_text_fallback():
    # Test that providing an unsupported file or missing OCR gracefully falls back
    # Just a mock file path that doesn't exist
    text = extract_text("dummy.xyz", "dummy.xyz")
    assert text == ""

def test_extract_xlsx(tmp_path):
    # Create a dummy xlsx using openpyxl
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = "Financial"
    ws['B1'] = "Report"
    filepath = tmp_path / "test.xlsx"
    wb.save(filepath)
    
    text = extract_text(filepath, "xlsx")
    assert "Financial" in text
    assert "Report" in text

def test_extract_eml(tmp_path):
    import email
    from email.message import EmailMessage
    msg = EmailMessage()
    msg.set_content("Please review the attached contract.")
    msg['Subject'] = "Contract Review"
    msg['From'] = "legal@office.com"
    msg['To'] = "hr@office.com"
    
    filepath = tmp_path / "test.eml"
    with open(filepath, "wb") as f:
        f.write(bytes(msg))
        
    text = extract_text(filepath, "eml")
    assert "Contract Review" in text
    assert "legal@office.com" in text
    assert "Please review the attached contract." in text
